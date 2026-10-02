import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Colors
    NAVY = RGBColor(15, 23, 42)       # #0F172A
    PRIMARY = RGBColor(0, 119, 182)    # #0077B6
    SKY = RGBColor(2, 132, 199)        # #0284C7
    SLATE = RGBColor(71, 85, 105)      # #475569
    LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC
    WHITE = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 232, 240)
    EMERALD = RGBColor(16, 185, 129)
    AMBER = RGBColor(245, 158, 11)

    screenshot_base = os.path.abspath("public/screenshots/presentation")

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text):
        # Badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.4))
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.55))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.4))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = SLATE

    def add_card(slide, left, top, width, height, bg_rgb=WHITE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_rgb
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: COVER
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1)

    # Big decorative card
    add_card(slide1, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))

    logo_path = os.path.abspath("public/logo.png")
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(5.916), Inches(1.5), width=Inches(1.5))

    tb_cover = slide1.shapes.add_textbox(Inches(2.0), Inches(3.2), Inches(9.333), Inches(2.8))
    tf_c = tb_cover.text_frame
    tf_c.word_wrap = True

    p0 = tf_c.paragraphs[0]
    p0.text = "SINTESA TPMPS"
    p0.alignment = PP_ALIGN.CENTER
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = PRIMARY

    p1 = tf_c.add_paragraph()
    p1.text = "SISTEM INFORMASI MANAJEMEN PENJAMINAN MUTU PENDIDIKAN SEKOLAH"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = NAVY

    p2 = tf_c.add_paragraph()
    p2.text = "SMK NEGERI 2 MAGELANG — TAHUN AJARAN 2025/2026\nSwadaya Bhina Raharja | Jl. Perintis Kemerdekaan No. 9, Kota Magelang"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(11)
    p2.font.color.rgb = SLATE

    p3 = tf_c.add_paragraph()
    p3.text = "\nPanduan Lengkap Hak Akses Pimpinan (Super Admin, Kepsek, Ketua TPMPS) & Rincian Tugas 15 Unit Kerja"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = SKY

    # =========================================================================
    # SLIDE 2: LOGIN PAGE & KEAMANAN SISTEM
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2)
    add_header(slide2, "MODUL AUTENTIKASI & KEAMANAN", "Halaman Login Terpusat & Kredensial Pengguna", "Gerbang masuk sistem dengan proteksi Supabase PostgreSQL RLS & Audit Logging Mutu")

    login_img = os.path.join(screenshot_base, "01_login_page.png")
    if os.path.exists(login_img):
        slide2.shapes.add_picture(login_img, Inches(0.8), Inches(1.9), width=Inches(7.2))

    card_login = add_card(slide2, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
    tb_l = slide2.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "FITUR & FUNGSI LOGIN:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY

    points_l = [
        ("Satu Akun Terintegrasi:", "Setiap pimpinan dan 15 unit kerja login menggunakan email kedinasan resmi (@smkn2magelang.sch.id) atau alias cepat."),
        ("Password Default:", "Seluruh akun awal diatur menggunakan kata sandi 'sintesa123' yang dapat diubah berkala."),
        ("Proteksi Peran Dinamis (RBAC):", "Sistem mengenali peran secara presisi: Super Admin, Kepala Sekolah, Ketua TPMPS, atau Akun Unit."),
        ("Session Cookie & LocalStorage:", "Menjamin persistensi sesi kerja tanpa terputus selama proses input evaluasi dan upload bukti dokumen."),
        ("Keamanan Multi-Lapis:", "Dilengkapi Row Level Security (RLS) pada basis data serta jejak forensik IP Address dan waktu aktivitas.")
    ]
    for title, desc in points_l:
        p = tf_l.add_paragraph()
        p.text = f"• {title} {desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 3: WEWENANG SUPER ADMIN, KEPSEK, & KETUA TPMPS
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3)
    add_header(slide3, "STRUKTUR OTORITAS TINGKAT PIMPINAN", "Wewenang Super Admin, Kepala Sekolah, & Ketua TPMPS", "Tiga pilar pimpinan penjaminan mutu pendidikan di SMK Negeri 2 Magelang")

    roles = [
        ("SUPER ADMIN", "Rian Prasetyo, S.Kom.", "admin.sintesa@smkn2magelang.sch.id", PRIMARY, [
            "CRUD Unit Kerja Penuh (/unit): Menambah unit baru, mengedit identitas/PIC, dan menghapus unit.",
            "Manajemen Akun: Mengelola kredensial dan hak akses seluruh akun sistem.",
            "Akses Penuh Google Drive Seluruh Unit: Membuka dan mengelola berkas 15 unit kerja.",
            "Konfigurasi Master SNP: Mengatur bobot 8 Standar Nasional Pendidikan.",
            "Audit Trail & Database Backup: Memantau log forensik aktivitas dan backup PostgreSQL."
        ]),
        ("KETUA TPMPS", "Dra. Hj. Siti Fatimah, M.M.", "tpmps.ketua@smkn2magelang.sch.id", SKY, [
            "Memimpin Siklus Mutu PPEPP: Mengawal penetapan standar hingga peningkatan mutu.",
            "Validasi & Verifikasi Bukti (/dokumen/validasi): Mengecek dan mengesahkan berkas unit.",
            "Audit Evaluasi Mandiri: Memberikan skor verifikasi auditor pada instrumen tiap unit.",
            "Persetujuan Program RTL: Menyetujui Rencana Tindak Lanjut dan pagu anggaran.",
            "Penerbitan Laporan EDS: Menyusun Rapor Mutu resmi untuk disahkan Kepala Sekolah."
        ]),
        ("KEPALA SEKOLAH", "Drs. H. Mulyono, M.Pd.", "kepala.sekolah@smkn2magelang.sch.id", NAVY, [
            "Executive Monitoring Dashboard (/dashboard): Memantau capaian agregat 8 SNP real-time.",
            "Approval Akhir Rapor Mutu: Mengesahkan secara legal Laporan EDS sekolah.",
            "Pengawasan Kepatuhan 15 Unit: Memantau statistik pemenuhan bukti fisik tiap unit.",
            "Otoritas CRUD Unit Kerja: Menyetujui penambahan atau mutasi unit kerja sekolah.",
            "Evaluasi Anggaran RTL: Menyetujui pagu anggaran pembenahan mutu sekolah."
        ])
    ]

    for idx, (role_name, pic_name, email_name, col, items) in enumerate(roles):
        left_pos = Inches(0.8 + idx * 4.0)
        add_card(slide3, left_pos, Inches(1.9), Inches(3.7), Inches(5.0))
        tb_r = slide3.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.1), Inches(3.3), Inches(4.6))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p = tf_r.paragraphs[0]
        p.text = role_name
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col

        p_pic = tf_r.add_paragraph()
        p_pic.text = f"{pic_name}\n{email_name}\n"
        p_pic.font.size = Pt(9)
        p_pic.font.color.rgb = SLATE

        for it in items:
            p_it = tf_r.add_paragraph()
            p_it.text = f"✔ {it}"
            p_it.font.size = Pt(9)
            p_it.font.color.rgb = NAVY

    # =========================================================================
    # SLIDE 4: DASHBOARD MUTU AGREGAT
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4)
    add_header(slide4, "MONITORING TINGKAT PIMPINAN", "Dashboard Penjaminan Mutu Agregat (/dashboard)", "Visualisasi capaian 8 Standar SNP, distribusi predikat unit, dan indikator kunci real-time")

    dash_img = os.path.join(screenshot_base, "02_dashboard_mutu.png")
    if os.path.exists(dash_img):
        slide4.shapes.add_picture(dash_img, Inches(0.8), Inches(1.9), width=Inches(7.2))

    add_card(slide4, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
    tb_d = slide4.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True

    p = tf_d.paragraphs[0]
    p.text = "YANG DAPAT DILAKUKAN PIMPINAN:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY

    pts_d = [
        ("Capaian 8 Standar SNP:", "Kepsek & TPMPS melihat skor agregat sekolah (misal: 87.2% Predikat Unggul) secara otomatis dari input 15 unit."),
        ("Grafik Tren Mutu Vokasi:", "Spline Chart memantau pergerakan mutu per semester dibanding target akreditasi sekolah (95.0%)."),
        ("Donut Distribusi Predikat:", "Memetakan sebaran predikat 15 unit kerja: Unggul (A), Baik (B), Cukup (C), atau Perlu Perhatian."),
        ("Kartu Metrik KPI Utama:", "Melihat langsung 8 Standar SNP, 26 Dokumen Terverifikasi, 5 Evaluasi Aktif, dan 2 Program RTL Berjalan."),
        ("Pusat Notifikasi:", "Badge notifikasi real-time memperingatkan jika ada berkas unit baru yang memerlukan review auditor TPMPS.")
    ]
    for title, desc in pts_d:
        p = tf_d.add_paragraph()
        p.text = f"• {title} {desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 5: VALIDASI DOKUMEN MUTU (TPMPS & KEPSEK MEMERIKSA UPLOAD UNIT)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5)
    add_header(slide5, "AUDIT & PENGAWASAN MUTU BUKTI", "Modul Validasi Dokumen Auditor (/dokumen/validasi)", "Fitur khusus TPMPS & Kepsek untuk mengecek, menelaah, merevisi, dan mengesahkan berkas tiap unit")

    val_img = os.path.join(screenshot_base, "03_validasi_dokumen_tpmps_kepsek.png")
    if os.path.exists(val_img):
        slide5.shapes.add_picture(val_img, Inches(0.8), Inches(1.9), width=Inches(7.2))

    add_card(slide5, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
    tb_v = slide5.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
    tf_v = tb_v.text_frame
    tf_v.word_wrap = True

    p = tf_v.paragraphs[0]
    p.text = "MEKANISME KONTROL UPLOAD UNIT:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY

    pts_v = [
        ("Mengecek Unggahan Tiap Unit:", "TPMPS dan Kepsek dapat melihat seluruh berkas yang diunggah oleh 15 unit kerja secara real-time di antrean terpusat."),
        ("Informasi Asal Unit & PIC:", "Setiap kartu menampilkan kode dokumen, nama unit pengunggah (misal: WKS 3, PPLG/RPL, Humas), nama berkas, ukuran, dan versi."),
        ("Pratinjau Dokumen (Preview):", "Pimpinan dapat membuka dan membaca berkas bukti langsung dari aplikasi tanpa perlu keluar browser."),
        ("Aksi 'Sahkan Dokumen':", "Jika berkas memenuhi standar, auditor mengklik Sahkan -> status berubah menjadi 'Terverifikasi' & skor unit otomatis bertambah."),
        ("Aksi 'Revisi' & 'Tolak':", "Jika bukti belum lengkap/kurang tepat, auditor memberikan catatan revisi yang langsung muncul di akun unit pengunggah.")
    ]
    for title, desc in pts_v:
        p = tf_v.add_paragraph()
        p.text = f"✔ {title} {desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 6: KELOLA UNIT KERJA (CRUD UNIT UNTUK SUPER ADMIN, KEPSEK, & TPMPS)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide6)
    add_header(slide6, "MANAJEMEN ORGANISASI SEKOLAH", "Modul Kelola Unit Kerja / CRUD Unit (/unit)", "Otoritas penuh Super Admin, Kepala Sekolah, dan Ketua TPMPS untuk mengelola struktur unit")

    crud_img = os.path.join(screenshot_base, "04_kelola_unit_crud.png")
    if os.path.exists(crud_img):
        slide6.shapes.add_picture(crud_img, Inches(0.8), Inches(1.9), width=Inches(7.2))

    add_card(slide6, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
    tb_c = slide6.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "FITUR CRUD UNIT KERJA:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY

    pts_c = [
        ("Create (+ Tambah Unit):", "Pimpinan dapat menambah unit kerja baru saat ada pembentukan pokja/satgas mutu baru lengkap dengan kode, PIC, email, dan SNP binaan."),
        ("Read (Pantau Kinerja Unit):", "Tabel ringkasan memuat 15 unit kerja, kategori (Manajemen, Kejuruan, Layanan), rasio indikator terpenuhi, dan skor capaian."),
        ("Update (Edit Unit):", "Mengubah pejabat PIC saat rotasi jabatan guru/tendik, memperbarui email resmi, atau menyesuaikan standar SNP binaan."),
        ("Delete (Hapus Unit):", "Menonaktifkan unit yang telah dilebur dengan dialog konfirmasi keselamatan dan pencatatan audit log forensik."),
        ("Proteksi Hak Akses:", "Hanya Super Admin, Kepsek, dan Ketua TPMPS yang memiliki tombol CRUD ini; akun guru/unit kerja hanya memiliki akses baca terbatas.")
    ]
    for title, desc in pts_c:
        p = tf_c.add_paragraph()
        p.text = f"✔ {title} {desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 7: BANK DOKUMEN DIGITAL (REPOSITORI ARSIP MUTU)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide7)
    add_header(slide7, "REPOSITORI DIGITAL TERPADU", "Bank Dokumen Mutu Digital Seluruh Unit (/dokumen)", "Pusat arsip bukti fisik seluruh 8 SNP dari 15 unit kerja sekolah")

    bank_img = os.path.join(screenshot_base, "05_bank_dokumen_digital.png")
    if os.path.exists(bank_img):
        slide7.shapes.add_picture(bank_img, Inches(0.8), Inches(1.9), width=Inches(7.2))

    add_card(slide7, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
    tb_b = slide7.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True

    p = tf_b.paragraphs[0]
    p.text = "FITUR BANK DOKUMEN:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY

    pts_b = [
        ("Arsip Lintas Unit:", "Menggabungkan bukti fisik dari seluruh 15 unit kerja dalam satu repositori yang rapi dan terstruktur."),
        ("Penyaringan Multi-Kriteria:", "Dapat difilter berdasarkan 8 Standar SNP, asal unit kerja, format berkas (PDF, Excel, Word), serta status validasi."),
        ("Pencarian Cerdas:", "Menemukan dokumen dalam hitungan detik menggunakan kata kunci judul, kode berkas, atau nama pengunggah."),
        ("Riwayat Verifikasi Lengkap:", "Mencantumkan tanggal validasi, nama auditor yang menyetujui, serta catatan verifikasi resmi."),
        ("Unduh & Pratinjau Cepat:", "Memudahkan saat visitasi asesor akreditasi BAN-SM atau audit ISO dengan akses dokumen satu pintu.")
    ]
    for title, desc in pts_b:
        p = tf_b.add_paragraph()
        p.text = f"• {title} {desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 8: ARSITEKTUR GOOGLE DRIVE PER UNIT KERJA (/drive)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide8)
    add_header(slide8, "REPOSITORI GOOGLE DRIVE UNIT", "Modul Google Drive Per Unit Kerja (/drive)", "Pengalaman repositori modern mirip Google Drive yang didedikasikan untuk 15 unit kerja")

    drive_overview_img = os.path.join(screenshot_base, "06_unit_01_wks1_kurikulum.png")
    if os.path.exists(drive_overview_img):
        slide8.shapes.add_picture(drive_overview_img, Inches(0.8), Inches(1.9), width=Inches(7.2))

    add_card(slide8, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
    tb_do = slide8.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
    tf_do = tb_do.text_frame
    tf_do.word_wrap = True

    p = tf_do.paragraphs[0]
    p.text = "FITUR UTAMA GOOGLE DRIVE UNIT:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY

    pts_do = [
        ("Daftar 15 Unit Kerja (Kiri):", "Panel kiri menampilkan 15 unit dengan nomor urut, kode unit, kategori, dan indikator unit yang sedang aktif."),
        ("10 Folder Standar SNP Baku:", "Tiap unit memiliki 10 folder baku (Standar SKL, Isi, Proses, Penilaian, Pendidik, Sarpras, Pengelolaan, Pembiayaan, SK, & Portofolio)."),
        ("Tombol '+ Baru' & 'Folder Baru':", "Unit dapat mengunggah berkas baru (PDF, Excel, Word, Foto) atau membuat sub-folder kustom."),
        ("Disarankan & Sering Diakses:", "Menampilkan kartu berkas terkini lengkap dengan ukuran, status verifikasi, dan pintasan akses cepat."),
        ("Indikator Kuota Penyimpanan:", "Visualisasi kapasitas penyimpanan unit (misal: 4.2 GB dari 15 GB terpakai).")
    ]
    for title, desc in pts_do:
        p = tf_do.add_paragraph()
        p.text = f"• {title} {desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = SLATE

    # =========================================================================
    # SLIDES 9 TO 23: 15 UNIT KERJA LENGKAP
    # =========================================================================
    units_data = [
        (
            1, "WKS 1 (Bidang Kurikulum)", "WKS-1", "Manajemen", "Dra. Sri Wahyuni, M.Pd.", "kurikulum@smkn2magelang.sch.id",
            "SNP-2 (Standar Isi), SNP-3 (Standar Proses), SNP-4 (Standar Penilaian)",
            "06_unit_01_wks1_kurikulum.png",
            [
                "Menyusun & mengunggah dokumen KOSP (Kurikulum Operasional Satuan Pendidikan) & Kurikulum Merdeka.",
                "Mengunggah Kalender Akademik Sekolah, SK Pembagian Beban Mengajar Guru, dan Jadwal Pelajaran.",
                "Mengisi evaluasi mandiri perangkat ajar (Silabus, Modul Ajar, ATP) dan instrumen asesmen diagnostik/sumatif.",
                "Mengajukan program RTL penyelarasan kurikulum bersama IDUKA mitra industri.",
                "Mengelola Google Drive Unit: Folder Kurikulum, Perangkat Pembelajaran, dan Bank Soal Ujian."
            ]
        ),
        (
            2, "WKS 2 (Bidang Kesiswaan)", "WKS-2", "Manajemen", "Bambang Sutrisno, S.Pd.", "kesiswaan@smkn2magelang.sch.id",
            "SNP-1 (Standar SKL), SNP-7 (Standar Pengelolaan)",
            "07_unit_02_wks2_kesiswaan.png",
            [
                "Mengunggah Buku Pedoman Tata Tertib Siswa, Buku Saku Karakter, dan SOP Penegakan Disiplin.",
                "Mendokumentasikan rekap prestasi lomba siswa di ajang LKS SMK, O2SN, FLS2N (Kota, Provinsi, Nasional).",
                "Mengarsipkan data kegiatan ekstrakurikuler, Pramuka Wajib, kepengurusan OSIS/MPK, dan beasiswa PIP/KIP.",
                "Mengisi evaluasi mandiri pembiasaan 7 Kebiasaan Anak Indonesia Hebat dan implementasi Projek P5.",
                "Mengelola Google Drive Unit: Folder Prestasi Siswa, Tata Tertib, dan Laporan Kesiswaan."
            ]
        ),
        (
            3, "WKS 3 (Bidang Sarana & Prasarana)", "WKS-3", "Manajemen", "Ir. Agus Haryanto, M.T.", "sarpras@smkn2magelang.sch.id",
            "SNP-6 (Standar Sarana & Prasarana)",
            "08_unit_03_wks3_sarpras.png",
            [
                "Mengunggah Buku Inventarisasi sarana prasarana sekolah, daftar peralatan bengkel, dan PC lab komputer.",
                "Menyusun Master Plan pemeliharaan gedung, ruang praktik siswa (RPS), instalasi listrik, dan air.",
                "Mengunggah SOP Keselamatan dan Kesehatan Kerja (K3) serta denah jalur evakuasi bencana sekolah.",
                "Mengajukan program RTL peremajaan workstation PC Lab RPL dan perbaikan ruang kelas.",
                "Mengelola Google Drive Unit: Folder Inventaris Aset, Pemeliharaan Gedung, dan SOP K3 Sarpras."
            ]
        ),
        (
            4, "WKS 4 (Hubungan Industri & Humas)", "WKS-4", "Manajemen", "Drs. Hendro Wibowo", "humas@smkn2magelang.sch.id",
            "SNP-2 (Standar Isi), SNP-3 (Standar Proses), SNP-7 (Standar Pengelolaan)",
            "09_unit_04_wks4_humas.png",
            [
                "Mengunggah naskah nota kesepahaman (MoU) dan Perjanjian Kerjasama (PKS) dengan DUDI mitra strategis.",
                "Mengarsipkan Berita Acara sinkronisasi kurikulum industri dan pelaksanaan Guru Tamu praktisi industri.",
                "Mengunggah laporan monitoring PKL siswa, jurnal pembimbingan industri, dan sertifikat PKL.",
                "Mengisi evaluasi kemitraan industri dan tingkat kepuasan mitra DUDI terhadap lulusan SMK.",
                "Mengelola Google Drive Unit: Folder MoU Industri, Laporan PKL, dan Dokumentasi Kunjungan Industri."
            ]
        ),
        (
            5, "Bidang Ketenagaan & SDM", "WKS-SDM", "Manajemen", "Nurul Hidayati, S.Pd.", "sdm@smkn2magelang.sch.id",
            "SNP-5 (Standar Pendidik & Tenaga Kependidikan)",
            "10_unit_05_wks_sdm.png",
            [
                "Mengunggah rekapitulasi kualifikasi akademik pendidik (S1/S2), Serdik, dan sertifikasi teknis kejuruan.",
                "Mengarsipkan data magang industri guru kejuruan (minimal 1 bulan) dan lisensi asesor kompetensi BNSP.",
                "Mengunggah rekapitulasi Penilaian Kinerja Guru (PKG) dan Sasaran Kinerja Pegawai (SKP) Tendik.",
                "Mengajukan program RTL peningkatan kompetensi pedagogik dan pelatihan kejuruan guru.",
                "Mengelola Google Drive Unit: Folder Sertifikasi Guru, Magang Industri Guru, dan Berkas PKG Tendik."
            ]
        ),
        (
            6, "Program Keahlian PPLG / RPL", "PROG-RPL", "Kejuruan", "Eko Prasetyo, S.Kom., M.Cs.", "rpl@smkn2magelang.sch.id",
            "SNP-1 (SKL), SNP-3 (Standar Proses), SNP-4 (Standar Penilaian)",
            "11_unit_06_prog_rpl.png",
            [
                "Mengunggah modul ajar produktif Coding Web, Mobile Development, Cloud Computing, dan Database.",
                "Mengunggah portofolio aplikasi perangkat lunak hasil Teaching Factory (TEFA) pesanan UMKM/klien riil.",
                "Mengunggah rekapitulasi penilaian Uji Kompetensi Keahlian (UKK) skema Rekayasa Perangkat Lunak.",
                "Mengarsipkan bukti sertifikasi kompetensi siswa dari Lembaga Sertifikasi Profesi (LSP-P1 / BNSP).",
                "Mengelola Google Drive Unit: Folder Modul Ajar RPL, Portofolio Proyek Aplikasi, dan Nilai UKK Siswa."
            ]
        ),
        (
            7, "Program Keahlian TJKT / TKJ", "PROG-TKJ", "Kejuruan", "Ahmad Fauzi, S.T.", "tkj@smkn2magelang.sch.id",
            "SNP-1 (SKL), SNP-3 (Standar Proses), SNP-6 (Standar Sarpras)",
            "12_unit_07_prog_tkj.png",
            [
                "Mengunggah modul ajar administrasi server Linux/Windows, keamanan siber, dan penyambungan fiber optik.",
                "Mengunggah diagram arsitektur jaringan backbone sekolah dan dokumentasi sertifikasi Mikrotik MTCNA / Cisco.",
                "Mengunggah Berita Acara dan rekapan nilai UKK Network Administrator bekerjasama dengan PT Telkom.",
                "Mengisi evaluasi kelayakan peralatan lab jaringan (Routerboard, Fusion Splicer, OTDR, Server).",
                "Mengelola Google Drive Unit: Folder Topologi Jaringan, Modul TKJ, dan Dokumentasi Sertifikasi Siswa."
            ]
        ),
        (
            8, "Program Keahlian Akuntansi (AKL)", "PROG-AKL", "Kejuruan", "Siti Rahmawati, S.E., M.Akt.", "akl@smkn2magelang.sch.id",
            "SNP-1 (SKL), SNP-2 (Standar Isi), SNP-4 (Standar Penilaian)",
            "13_unit_08_prog_akl.png",
            [
                "Mengunggah lembar kerja siklus akuntansi perusahaan jasa, dagang, dan manufaktur (MYOB & Accurate).",
                "Mengunggah modul praktikum perpajakan elektronik (e-Faktur, e-SPT) dan pembukuan Mini Bank sekolah.",
                "Mengunggah rekapan hasil sertifikasi Teknisi Akuntansi Yunior berlisensi LSP-P1 dan BNSP.",
                "Mendokumentasikan kegiatan simulasi audit pembukuan keuangan dan rekonsiliasi kas.",
                "Mengelola Google Drive Unit: Folder Modul Akuntansi Komputer, Laporan Praktik Mini Bank, dan Berkas UKK."
            ]
        ),
        (
            9, "Program Keahlian Manajemen Perkantoran (MPLB)", "PROG-OTKP", "Kejuruan", "Dewi Lestari, S.Pd.", "mplb@smkn2magelang.sch.id",
            "SNP-1 (SKL), SNP-3 (Standar Proses), SNP-7 (Standar Pengelolaan)",
            "14_unit_09_prog_mplb.png",
            [
                "Mengunggah SOP Manajemen Kearsipan Digital (e-filing), tata persuratan, dan korespondensi bisnis.",
                "Mengunggah portofolio simulasi rapat bisnis, pelayanan prima (service excellence), dan keprotokolan humas.",
                "Mengunggah rekapitulasi penilaian UKK skema Otomasi Tata Kelola Perkantoran tersertifikasi BNSP.",
                "Mengisi evaluasi kesiapan sarana lab perkantoran (mesin fax, filling cabinet, shredder, PC sekretaris).",
                "Mengelola Google Drive Unit: Folder SOP Kearsipan, Modul MPLB, dan Portofolio Praktik Humas Kantor."
            ]
        ),
        (
            10, "Program Keahlian Pemasaran (BDP)", "PROG-BDP", "Kejuruan", "Rudi Hartono, S.E.", "bdp@smkn2magelang.sch.id",
            "SNP-1 (SKL), SNP-3 (Standar Proses), SNP-4 (Standar Penilaian)",
            "15_unit_10_prog_bdp.png",
            [
                "Mengunggah laporan omzet praktik penjualan retail di Business Center dan toko vokasi sekolah.",
                "Mengunggah portofolio kampanye digital marketing siswa (Social Media Marketing, SEO, TikTok Shop).",
                "Mengunggah bukti uji sertifikasi kompetensi pramuniaga dan kasir berlisensi LSP/BNSP.",
                "Mengarsipkan modul ajar visual merchandising, penataan produk, dan negosiasi bisnis.",
                "Mengelola Google Drive Unit: Folder Laporan Omzet Retail, Portofolio Digital Marketing, dan Uji Kompetensi."
            ]
        ),
        (
            11, "Unit Pengelola Bengkel & Lab Komputer", "BENGKEL", "Layanan", "Supriyanto, A.Md.", "lab@smkn2magelang.sch.id",
            "SNP-6 (Standar Sarpras), SNP-7 (Standar Pengelolaan)",
            "16_unit_11_lab_bengkel.png",
            [
                "Mengunggah jadwal penggunaan harian seluruh laboratorium komputer dan bengkel kejuruan sekolah.",
                "Mengunggah kartu riwayat perawatan dan perbaikan hardware (maintenance log) PC, server, dan printer.",
                "Mengisi evaluasi mandiri kesiapan PC dan jaringan untuk pelaksanaan Computer-Based Test (CBT 1:1).",
                "Mengarsipkan formulir peminjaman alat bengkel dan SOP keselamatan kerja teknisi laboratorium.",
                "Mengelola Google Drive Unit: Folder Jadwal Penggunaan Lab, Kartu Maintenance PC, dan SOP Bengkel."
            ]
        ),
        (
            12, "Unit Perpustakaan Digital", "PERPUS", "Layanan", "Tri Utami, S.I.Pust.", "perpustakaan@smkn2magelang.sch.id",
            "SNP-3 (Standar Proses), SNP-6 (Standar Sarpras)",
            "17_unit_12_perpustakaan.png",
            [
                "Mengunggah katalog buku digital (OPAC) dan rekap statistik kunjungan peminjaman e-library siswa/guru.",
                "Mengunggah program pembiasaan literasi sekolah dan naskah MoU kerjasama dengan Perpustakaan Daerah.",
                "Mengarsipkan daftar pengadaan buku teks Kurikulum Merdeka dan buku referensi kejuruan terbaru.",
                "Mengisi evaluasi mandiri kecukupan judul buku dan kenyamanan fasilitas ruang baca digital.",
                "Mengelola Google Drive Unit: Folder Katalog Buku OPAC, Statistik Kunjungan, dan Berkas Literasi Vokasi."
            ]
        ),
        (
            13, "Bursa Kerja Khusus (BKK Magelang)", "BKK", "Layanan", "Wahyu Nugroho, S.Pd.", "bkk@smkn2magelang.sch.id",
            "SNP-1 (Standar SKL), SNP-7 (Standar Pengelolaan)",
            "18_unit_13_bkk.png",
            [
                "Mengunggah laporan Tracer Study keterserapan alumni (Bekerja, Melanjutkan Kuliah, Wirausaha - BMW).",
                "Mengunggah dokumentasi pelaksanaan Campus Recruitment dan Job Fair SMK Negeri 2 Magelang.",
                "Mengunggah statistik penyaluran tenaga kerja alumni ke industri nasional dan program magang ke Jepang/Jerman.",
                "Mengisi evaluasi masa tunggu lulusan hingga mendapatkan pekerjaan pertama (target < 3 bulan).",
                "Mengelola Google Drive Unit: Folder Laporan Tracer Study BMW, Dokumentasi Job Fair, dan Rekap Penyaluran Alumni."
            ]
        ),
        (
            14, "Unit Bimbingan Konseling (BK)", "BK", "Layanan", "Dra. Endang Sulastri", "bk@smkn2magelang.sch.id",
            "SNP-3 (Standar Proses), SNP-7 (Standar Pengelolaan)",
            "19_unit_14_bk.png",
            [
                "Mengunggah Program Tahunan Layanan BK (Bimbingan Pribadi, Sosial, Belajar, dan Perencanaan Karir).",
                "Mengunggah instrumen asesmen diagnostik non-kognitif pemetaan gaya belajar dan minat karir siswa baru.",
                "Mengunggah dokumentasi program pencegahan perundungan (anti-bullying) dan penciptaan sekolah ramah anak.",
                "Mengarsipkan laporan rekapitulasi penanganan konseling individual dan bimbingan kelompok.",
                "Mengelola Google Drive Unit: Folder Program Tahunan BK, Asesmen Diagnostik, dan Laporan Konseling Siswa."
            ]
        ),
        (
            15, "Unit Produksi & Teaching Factory (TEFA)", "TEFA", "Kejuruan", "Anwar Sadat, S.T.", "tefa@smkn2magelang.sch.id",
            "SNP-3 (Standar Proses), SNP-8 (Standar Pembiayaan)",
            "20_unit_15_tefa.png",
            [
                "Mengunggah SOP operasional Teaching Factory (TEFA) berstandar tata kelola industri manufaktur/jasa.",
                "Mengunggah neraca laporan omzet produksi barang/jasa vokasi dan kontribusi pendapatan ke kas BLUD sekolah.",
                "Mengunggah portofolio produk unggulan, katalog layanan jasa kejuruan, dan nota pesanan order klien industri.",
                "Mengisi evaluasi keterlibatan siswa dalam siklus produksi riil berorientasi pasar konsumen.",
                "Mengelola Google Drive Unit: Folder Neraca Keuangan TEFA, Katalog Produk Unggulan, dan SOP Produksi Vokasi."
            ]
        )
    ]

    for unit in units_data:
        num, name, code, cat, pic, email, snp, img_file, tasks = unit
        slide_u = prs.slides.add_slide(blank_slide_layout)
        set_slide_background(slide_u)
        badge_u = f"UNIT KERJA {num} &bull; KATEGORI {cat.upper()}"
        title_u = f"{name} ({code})"
        sub_u = f"PIC: {pic} | Email: {email} | Diampu: {snp}"
        add_header(slide_u, badge_u, title_u, sub_u)

        # Left: Full screenshot
        img_path = os.path.join(screenshot_base, img_file)
        if os.path.exists(img_path):
            slide_u.shapes.add_picture(img_path, Inches(0.8), Inches(1.9), width=Inches(7.2))

        # Right: Tasks & Functionality Card
        add_card(slide_u, Inches(8.3), Inches(1.9), Inches(4.2), Inches(5.0))
        tb_u = slide_u.shapes.add_textbox(Inches(8.5), Inches(2.1), Inches(3.8), Inches(4.6))
        tf_u = tb_u.text_frame
        tf_u.word_wrap = True

        p = tf_u.paragraphs[0]
        p.text = "TUGAS POKOK & BISA APA DI SISTEM:"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = NAVY

        p_snp = tf_u.add_paragraph()
        p_snp.text = f"Standar SNP: {snp}\n"
        p_snp.font.size = Pt(9)
        p_snp.font.bold = True
        p_snp.font.color.rgb = PRIMARY

        for t in tasks:
            p_t = tf_u.add_paragraph()
            p_t.text = f"✔ {t}"
            p_t.font.size = Pt(8.8)
            p_t.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 24: ALUR VALIDASI DOKUMEN MUTU (END-TO-END WORKFLOW)
    # =========================================================================
    slide24 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide24)
    add_header(slide24, "ALUR KERJA PENJAMINAN MUTU", "Alur Kerja Terpadu: Dari Upload Unit ke Validasi Pimpinan", "Bagaimana integrasi data terjadi dari 15 unit kerja hingga evaluasi pimpinan sekolah")

    steps_wf = [
        ("01. Upload Unit Kerja", "15 Unit Kerja mengunggah berkas bukti fisik (PDF/Excel/Word) ke folder 10 SNP di Google Drive Unit masing-masing.", PRIMARY),
        ("02. Masuk Antrean Review", "Sistem otomatis mendaftarkan berkas dengan status 'Menunggu Review' di modul /dokumen/validasi auditor TPMPS.", SKY),
        ("03. Pemeriksaan Auditor TPMPS", "Ketua & Tim TPMPS menelaah keabsahan dokumen bukti fisik melalui fitur Pratinjau Dokumen.", AMBER),
        ("04. Keputusan & Catatan Revisi", "Auditor mengklik 'Sahkan Dokumen' untuk berkas valid, atau 'Revisi' dengan catatan perbaikan untuk unit kerja.", EMERALD),
        ("05. Rekapitulasi Rapor Mutu", "Skor ketercapaian 8 SNP dan Rapor Mutu sekolah di /dashboard otomatis terupdate secara agregat untuk disahkan Kepala Sekolah.", NAVY)
    ]

    for idx, (st_title, st_desc, st_col) in enumerate(steps_wf):
        top_pos = Inches(1.9 + idx * 1.0)
        card_wf = add_card(slide24, Inches(0.8), top_pos, Inches(11.733), Inches(0.85))
        tb_wf = slide24.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.1), Inches(11.2), Inches(0.65))
        tf_wf = tb_wf.text_frame
        tf_wf.word_wrap = True

        p = tf_wf.paragraphs[0]
        p.text = st_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = st_col

        p_d = tf_wf.add_paragraph()
        p_d.text = st_desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = SLATE

    # =========================================================================
    # SLIDE 25: PENUTUP & PENGESAHAN MUTU
    # =========================================================================
    slide25 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide25)
    add_header(slide25, "KOMITMEN MUTU BERKELANJUTAN", "Pengesahan & Komitmen Mutu Bersama SPMI", "SINTESA TPMPS &bull; SMK Negeri 2 Magelang &bull; Swadaya Bhina Raharja")

    # Card 1: Kepala Sekolah
    add_card(slide25, Inches(1.5), Inches(2.2), Inches(4.8), Inches(4.5))
    tb_ks = slide25.shapes.add_textbox(Inches(1.8), Inches(2.5), Inches(4.2), Inches(4.0))
    tf_ks = tb_ks.text_frame
    tf_ks.word_wrap = True

    p = tf_ks.paragraphs[0]
    p.text = "Mengetahui,\nKepala SMK Negeri 2 Magelang"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE

    p_sp = tf_ks.add_paragraph()
    p_sp.text = "\n\n( Tanda Tangan & Cap Resmi )\n\n"
    p_sp.alignment = PP_ALIGN.CENTER
    p_sp.font.size = Pt(10)
    p_sp.font.color.rgb = RGBColor(148, 163, 184)

    p_nm = tf_ks.add_paragraph()
    p_nm.text = "Drs. H. Mulyono, M.Pd."
    p_nm.alignment = PP_ALIGN.CENTER
    p_nm.font.size = Pt(14)
    p_nm.font.bold = True
    p_nm.font.color.rgb = NAVY

    p_nip = tf_ks.add_paragraph()
    p_nip.text = "NIP. 19680312 199203 1 004"
    p_nip.alignment = PP_ALIGN.CENTER
    p_nip.font.size = Pt(10)
    p_nip.font.color.rgb = SLATE

    # Card 2: Ketua TPMPS
    add_card(slide25, Inches(7.0), Inches(2.2), Inches(4.8), Inches(4.5))
    tb_tp = slide25.shapes.add_textbox(Inches(7.3), Inches(2.5), Inches(4.2), Inches(4.0))
    tf_tp = tb_tp.text_frame
    tf_tp.word_wrap = True

    p = tf_tp.paragraphs[0]
    p.text = "Disusun Oleh,\nKetua TPMPS SMK Negeri 2 Magelang"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE

    p_sp2 = tf_tp.add_paragraph()
    p_sp2.text = "\n\n( Tanda Tangan Auditor )\n\n"
    p_sp2.alignment = PP_ALIGN.CENTER
    p_sp2.font.size = Pt(10)
    p_sp2.font.color.rgb = RGBColor(148, 163, 184)

    p_nm2 = tf_tp.add_paragraph()
    p_nm2.text = "Dra. Hj. Siti Fatimah, M.M."
    p_nm2.alignment = PP_ALIGN.CENTER
    p_nm2.font.size = Pt(14)
    p_nm2.font.bold = True
    p_nm2.font.color.rgb = NAVY

    p_nip2 = tf_tp.add_paragraph()
    p_nip2.text = "NIP. 19750918 200212 2 001"
    p_nip2.alignment = PP_ALIGN.CENTER
    p_nip2.font.size = Pt(10)
    p_nip2.font.color.rgb = SLATE

    # Save presentations
    output_path1 = os.path.abspath("PRESENTASI_SINTESA_TPMPS_LENGKAP.pptx")
    output_path2 = os.path.abspath("public/PRESENTASI_SINTESA_TPMPS_LENGKAP.pptx")
    prs.save(output_path1)
    prs.save(output_path2)
    print(f"Presentation saved successfully to:\n1) {output_path1}\n2) {output_path2}")

if __name__ == "__main__":
    create_presentation()
