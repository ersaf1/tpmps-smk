import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Color Palette
C_NAVY = RGBColor(15, 23, 42)
C_BLUE = RGBColor(37, 99, 235)
C_MUTED = RGBColor(100, 116, 139)
C_DARK = RGBColor(30, 41, 59)
C_AMBER = RGBColor(217, 119, 6)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Title & Cover Header
p_org = doc.add_paragraph()
p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_org = p_org.add_run("PEMERINTAH PROVINSI JAWA TENGAH • DINAS PENDIDIKAN DAN KEBUDAYAAN\nSMK NEGERI 2 MAGELANG — SWADAYA BHINA RAHARJA")
r_org.font.name = "Segoe UI"
r_org.font.size = Pt(9.5)
r_org.font.bold = True
r_org.font.color.rgb = C_MUTED

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t = p_title.add_run("BUKU PANDUAN LENGKAP PENGOPERASIAN SISTEM\nSINTESA TPMPS")
r_t.font.name = "Segoe UI"
r_t.font.size = Pt(22)
r_t.font.bold = True
r_t.font.color.rgb = C_NAVY

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_s = p_sub.add_run("Panduan Teknis Langkah demi Langkah (Step-by-Step User Manual)\nUntuk 15 Unit Kerja, Kepala Sekolah, Ketua TPMPS, dan Administrator")
r_s.font.name = "Segoe UI"
r_s.font.size = Pt(11)
r_s.font.italic = True
r_s.font.color.rgb = C_BLUE

doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_chapter(title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(title)
    r.font.name = "Segoe UI"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = C_BLUE

def add_subheading(title):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(title)
    r.font.name = "Segoe UI"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = C_NAVY

def add_body(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix + " ")
        rb.font.name = "Segoe UI"
        rb.font.size = Pt(10)
        rb.font.bold = True
        rb.font.color.rgb = C_NAVY
    rt = p.add_run(text)
    rt.font.name = "Segoe UI"
    rt.font.size = Pt(10)
    rt.font.italic = italic
    rt.font.color.rgb = C_DARK
    return p

def add_bullet(bold_label, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    rb = p.add_run(bold_label + ": ")
    rb.font.name = "Segoe UI"
    rb.font.size = Pt(10)
    rb.font.bold = True
    rb.font.color.rgb = C_NAVY
    rt = p.add_run(text)
    rt.font.name = "Segoe UI"
    rt.font.size = Pt(10)
    rt.font.color.rgb = C_DARK

# Read and parse PANDUAN_LENGKAP_SINTESA_TPMPS.md to convert into docx
md_path = r"C:\Users\lulus\tpmps-smk\PANDUAN_LENGKAP_SINTESA_TPMPS.md"
with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

current_section = ""
in_table = False
table_lines = []

for line in lines:
    line_str = line.strip()
    if not line_str:
        continue
    
    # Skip Title / Cover that we already handled
    if line_str.startswith("# BUKU PANDUAN") or line_str.startswith("## PANDUAN TEKNIS") or line_str.startswith("### SISTEM INFORMASI"):
        continue
    if line_str == "---":
        continue
    if line_str.startswith("## DAFTAR ISI"):
        continue
    if line_str.startswith("1. **BAB") or line_str.startswith("2. **BAB") or line_str.startswith("3. **BAB") or line_str.startswith("4. **BAB") or line_str.startswith("5. **BAB") or line_str.startswith("6. **BAB") or line_str.startswith("7. **BAB"):
        continue
    if line_str.startswith("- 1.") or line_str.startswith("- 2.") or line_str.startswith("- 3.") or line_str.startswith("- 4.") or line_str.startswith("- 5.") or line_str.startswith("- 6.") or line_str.startswith("- 7."):
        continue

    # Table handling
    if line_str.startswith("|"):
        if "---" in line_str:
            continue
        parts = [p.strip() for p in line_str.split("|")[1:-1]]
        table_lines.append(parts)
        in_table = True
        continue
    else:
        if in_table and table_lines:
            # Render Table
            cols_count = len(table_lines[0])
            t = doc.add_table(rows=len(table_lines), cols=cols_count)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, r_data in enumerate(table_lines):
                for c_idx, c_val in enumerate(r_data):
                    if c_idx < cols_count:
                        cell = t.cell(r_idx, c_idx)
                        cell.text = c_val.replace("**", "")
                        cp = cell.paragraphs[0]
                        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
                        cp.runs[0].font.name = "Segoe UI"
                        cp.runs[0].font.size = Pt(8.5)
                        if r_idx == 0:
                            cp.runs[0].font.bold = True
                            set_cell_background(cell, "0F172A")
                            cp.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F8FAFC")
                            cp.runs[0].font.color.rgb = C_DARK
                        set_cell_margins(cell)
            doc.add_paragraph().paragraph_format.space_after = Pt(8)
            table_lines = []
            in_table = False

    # Chapter Header #
    if line_str.startswith("# BAB"):
        add_chapter(line_str.replace("#", "").strip())
    # Subheading ###
    elif line_str.startswith("### "):
        add_subheading(line_str.replace("###", "").strip())
    # Sub-subheading ####
    elif line_str.startswith("#### "):
        add_body(line_str.replace("####", "").strip(), bold_prefix="▶")
    # Bullet points *
    elif line_str.startswith("* **") or line_str.startswith("- **"):
        parts = line_str.split("**")
        if len(parts) >= 3:
            bold_label = parts[1].strip()
            rest = "**".join(parts[2:]).lstrip(":").strip()
            add_bullet(bold_label, rest)
        else:
            add_body(line_str.replace("*", "").replace("-", "").strip())
    elif line_str.startswith("* ") or line_str.startswith("- "):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(line_str[2:].strip().replace("**", ""))
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.color.rgb = C_DARK
    elif line_str.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(line_str.replace("> ", "").replace("**", "").strip())
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = C_AMBER
    else:
        add_body(line_str.replace("**", ""))

# Save to output locations
out_file1 = r"C:\Users\lulus\tpmps-smk\PANDUAN_LENGKAP_SINTESA_TPMPS.docx"
out_file2 = r"C:\Users\lulus\Downloads\PANDUAN_LENGKAP_SINTESA_TPMPS.docx"

doc.save(out_file1)
doc.save(out_file2)
print(f"Comprehensive Word manual successfully generated at:\n1. {out_file1}\n2. {out_file2}")
