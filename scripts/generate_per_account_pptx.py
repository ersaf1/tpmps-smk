import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Color Palette Definitions
C_NAVY_DARK = RGBColor(15, 23, 42)      # #0f172a
C_NAVY_MED  = RGBColor(30, 41, 59)      # #1e293b
C_BLUE_MAIN = RGBColor(37, 99, 235)     # #2563eb
C_BLUE_LIGHT = RGBColor(239, 246, 255)  # #eff6ff
C_BLUE_BORDER = RGBColor(191, 219, 254) # #bfdbfe
C_TEXT_DARK = RGBColor(15, 23, 42)      # #0f172a
C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748b
C_TEXT_LIGHT = RGBColor(241, 245, 249)  # #f1f5f9
C_WHITE     = RGBColor(255, 255, 255)
C_AMBER     = RGBColor(217, 119, 6)     # #d97706
C_AMBER_BG  = RGBColor(254, 243, 199)   # #fef3c7
C_EMERALD   = RGBColor(16, 185, 129)    # #10b981
C_CARD_BG   = RGBColor(248, 250, 252)   # #f8fafc
C_CARD_BORDER = RGBColor(226, 232, 240) # #e2e8f0

SCREENSHOT_DIR = r"C:\Users\lulus\tpmps-smk\public\screenshots\exact_accounts"

def add_header(slide, badge_text, title_text, subtitle_text="Sistem Informasi Penjaminan Mutu Internal - SMK Negeri 2 Magelang"):
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    # Badge
    p0 = tf.paragraphs[0]
    p0.text = badge_text.upper()
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE_MAIN
    p0.space_after = Pt(2)
    
    # Title
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY_DARK
    p1.space_after = Pt(2)
    
    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_TEXT_MUTED
    
    # Accent Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.55), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = C_CARD_BORDER
    line.line.color.rgb = C_CARD_BORDER

def add_card_box(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def add_screenshot(slide, img_path, left, top, width, height):
    add_card_box(slide, left - Inches(0.08), top - Inches(0.08), width + Inches(0.16), height + Inches(0.16), bg_color=C_WHITE, border_color=C_CARD_BORDER)
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, left, top, width, height)
    else:
        ph = slide.shapes.add_textbox(left, top + Inches(1.5), width, Inches(1))
        p = ph.text_frame.paragraphs[0]
        p.text = f"[Screenshot: {os.path.basename(img_path)}]"
        p.alignment = PP_ALIGN.CENTER
        p.font.color.rgb = C_TEXT_MUTED

# ==========================================
# SLIDE 1: COVER
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY_DARK
bg1.line.fill.background()

bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.5), Inches(0.15), Inches(4.5))
bar1.fill.solid()
bar1.fill.fore_color.rgb = C_BLUE_MAIN
bar1.line.fill.background()

c_box = slide1.shapes.add_textbox(Inches(1.6), Inches(1.5), Inches(10.5), Inches(4.5))
tf1 = c_box.text_frame
tf1.word_wrap = True

p_badge = tf1.paragraphs[0]
p_badge.text = "SISTEM INFORMASI PENJAMINAN MUTU INTERNAL (SPMI) - SINTESA TPMPS"
p_badge.font.name = "Segoe UI"
p_badge.font.size = Pt(13)
p_badge.font.bold = True
p_badge.font.color.rgb = RGBColor(96, 165, 250)
p_badge.space_after = Pt(12)

p_t1 = tf1.add_paragraph()
p_t1.text = "SINTESA TPMPS"
p_t1.font.name = "Segoe UI"
p_t1.font.size = Pt(40)
p_t1.font.bold = True
p_t1.font.color.rgb = C_WHITE
p_t1.space_after = Pt(4)

p_t2 = tf1.add_paragraph()
p_t2.text = "Sistem Integrasi Penjaminan Mutu Terpadu SMK Negeri 2 Magelang"
p_t2.font.name = "Segoe UI"
p_t2.font.size = Pt(22)
p_t2.font.color.rgb = RGBColor(226, 232, 240)
p_t2.space_after = Pt(20)

p_desc = tf1.add_paragraph()
p_desc.text = "Presentasi Hasil Screenshot Asli Setiap Akun Unit Kerja (Bukan dari Super Admin),\nFitur Login 1-Klik Per Unit, Portal Pemeriksaan Berkas Baru, & Portofolio 15 Unit Resmi."
p_desc.font.name = "Segoe UI"
p_desc.font.size = Pt(14)
p_desc.font.color.rgb = RGBColor(148, 163, 184)
p_desc.space_after = Pt(36)

p_meta = tf1.add_paragraph()
p_meta.text = "SMK NEGERI 2 MAGELANG  •  SWADAYA BHINA RAHARJA  •  TAHUN AJARAN 2026/2027"
p_meta.font.name = "Segoe UI"
p_meta.font.size = Pt(11)
p_meta.font.bold = True
p_meta.font.color.rgb = C_AMBER

# ==========================================
# SLIDE 2: LOGIN & FITUR LOGIN PER UNIT
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "Modul 00: Autentikasi Pengguna", "Halaman Login & Fitur 1-Klik Login Per Unit")

card_l2 = add_card_box(slide2, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.1))
tfl2 = card_l2.text_frame
tfl2.word_wrap = True
tfl2.margin_left = tfl2.margin_right = tfl2.margin_top = Inches(0.25)

p = tfl2.paragraphs[0]
p.text = "FITUR LOGIN PER UNIT"
p.font.name = "Segoe UI"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = C_BLUE_MAIN
p.space_after = Pt(10)

features_login = [
    ("Login Cepat 1-Klik Per Unit", "Disediakan tombol khusus untuk langsung login ke masing-masing dari 15 unit resmi, KASEK, dan TPMPS tanpa harus mengetik."),
    ("Sesi & Identitas Terkunci", "Setiap akun memiliki sesi terpisah: WKS 1 masuk ke ruang kerja WKS 1, Kesiswaan ke ruang kesiswaan, dsb."),
    ("Otoritas Berbasis Peran (RBAC)", "Pimpinan memiliki hak validasi & pengesahan, sedangkan 14 unit pelaksana memiliki hak unggah dan pengelolaan berkas drive."),
    ("Dukungan Kode Singkat", "Pengguna juga dapat mengetik kode singkat di input login: wks1, wks2, k3, renbang, katu, kalab, naswil, bk, bkk, ups, kasek, tpmps.")
]

for title, desc in features_login:
    pt = tfl2.add_paragraph()
    pt.text = f"• {title}"
    pt.font.name = "Segoe UI"
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = C_NAVY_DARK
    
    pd = tfl2.add_paragraph()
    pd.text = f"  {desc}"
    pd.font.name = "Segoe UI"
    pd.font.size = Pt(10)
    pd.font.color.rgb = C_TEXT_MUTED
    pd.space_after = Pt(8)

add_screenshot(slide2, os.path.join(SCREENSHOT_DIR, "00_login", "01_login_page.png"), Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.1))

# ==========================================
# SLIDE 3: WEWENANG PIMPINAN & SUPER ADMIN
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "Struktur Kewenangan", "Matriks Peran: Super Admin, Kepala Sekolah (1. KASEK), & Ketua TPMPS")

roles_data = [
    ("SUPER ADMIN", "Pengendali Sistem & Konfigurasi", C_NAVY_MED, [
        "Kelola Master 15 Unit Kerja: Mengatur identitas, PIC, dan kategori",
        "Manajemen Akun Pengguna: Generate akun & reset password unit",
        "Konfigurasi Butir SNP: Menyesuaikan indikator mutu & instrumen",
        "Backup & Audit Log: Menjaga integritas dan keutuhan database"
    ]),
    ("1. KASEK (KEPALA SEKOLAH)", "Pengambil Keputusan & Pengesah Mutu", C_BLUE_MAIN, [
        "Executive Dashboard: Memantau ketercapaian 8 SNP sekolah",
        "Validasi Akhir Dokumen: Menyetujui / menolak berkas dari 15 unit",
        "Pengesahan Laporan EDS: Memberikan legalitas laporan formal",
        "Monitoring Rencana Tindak Lanjut: Memastikan eksekusi perbaikan"
    ]),
    ("7. KETUA TPMPS", "Koordinator & Verifikator Mutu", C_AMBER, [
        "Audit & Verifikasi Bukti Fisik: Menilai kesesuaian dokumen unit",
        "Input Skor Evaluasi Mutu: Menilai kepatuhan butir instrumen SNP",
        "Umpan Balik / Revisi: Memberi catatan langsung ke unit bersangkutan",
        "Penyusunan RTL: Merancang aksi mutu & penetapan target waktu"
    ])
]

for idx, (r_title, r_sub, r_color, r_bullets) in enumerate(roles_data):
    x_pos = Inches(0.8 + idx * 4.0)
    card_r = add_card_box(slide3, x_pos, Inches(1.8), Inches(3.7), Inches(5.1))
    tfr = card_r.text_frame
    tfr.word_wrap = True
    tfr.margin_left = Inches(0.2)
    tfr.margin_right = Inches(0.2)
    tfr.margin_top = Inches(0.25)
    
    pt = tfr.paragraphs[0]
    pt.text = r_title
    pt.font.name = "Segoe UI"
    pt.font.size = Pt(15)
    pt.font.bold = True
    pt.font.color.rgb = r_color
    pt.space_after = Pt(2)
    
    ps = tfr.add_paragraph()
    ps.text = r_sub
    ps.font.name = "Segoe UI"
    ps.font.size = Pt(9.5)
    ps.font.italic = True
    ps.font.color.rgb = C_TEXT_MUTED
    ps.space_after = Pt(16)
    
    for b in r_bullets:
        pb = tfr.add_paragraph()
        pb.text = f"✓ {b}"
        pb.font.name = "Segoe UI"
        pb.font.size = Pt(10)
        pb.font.color.rgb = C_NAVY_DARK
        pb.space_after = Pt(8)

# ==========================================
# SLIDE 4: DASHBOARD EKSEKUTIF KEPSEK (AKUN KASEK)
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "Pimpinan: Akun 1. KASEK", "Tampilan Dashboard Mutu Eksekutif (Login sebagai Kepala Sekolah)")

card_l4 = add_card_box(slide4, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.1))
tfl4 = card_l4.text_frame
tfl4.word_wrap = True
tfl4.margin_left = tfl4.margin_right = tfl4.margin_top = Inches(0.25)

p = tfl4.paragraphs[0]
p.text = "FITUR DASHBOARD KASEK"
p.font.name = "Segoe UI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_BLUE_MAIN
p.space_after = Pt(10)

kepsek_items = [
    ("Monitoring Capaian 8 SNP", "Grafik radar dan progress capaian Standar Nasional Pendidikan secara real-time."),
    ("Status Kelengkapan 15 Unit", "Melihat unit mana yang sudah mengunggah berkas lengkap, pending, atau memerlukan revisi."),
    ("Notifikasi Dokumen Masuk", "Indikator berkas baru yang menunggu persetujuan dan pengesahan Kepala Sekolah."),
    ("Sesi Resmi Kepala Sekolah", "Tampilan akun resmi Drs. H. Mulyono, M.Pd. dengan otorisasi pengesahan tertinggi.")
]

for title, desc in kepsek_items:
    pt = tfl4.add_paragraph()
    pt.text = f"• {title}"
    pt.font.name = "Segoe UI"
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = C_NAVY_DARK
    
    pd = tfl4.add_paragraph()
    pd.text = f"  {desc}"
    pd.font.name = "Segoe UI"
    pd.font.size = Pt(10)
    pd.font.color.rgb = C_TEXT_MUTED
    pd.space_after = Pt(8)

add_screenshot(slide4, os.path.join(SCREENSHOT_DIR, "01_kasek", "01_dashboard_kasek.png"), Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.1))

# ==========================================
# SLIDE 5: MODUL VALIDASI DOKUMEN (TAMPILAN BARU EKSEKUTIF)
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "Pimpinan: Validasi Dokumen", "Portal Baru Pemeriksaan Berkas Unggahan 15 Unit Kerja")

card_l5 = add_card_box(slide5, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.1))
tfl5 = card_l5.text_frame
tfl5.word_wrap = True
tfl5.margin_left = tfl5.margin_right = tfl5.margin_top = Inches(0.25)

p = tfl5.paragraphs[0]
p.text = "TAMPILAN BARU PEMERIKSAAN DOKUMEN"
p.font.name = "Segoe UI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_AMBER
p.space_after = Pt(8)

validasi_steps = [
    ("1. 4 Kartu KPI Mutu", "Menampilkan statistik Total Berkas, Menunggu Validasi, Telah Disahkan, dan Perlu Revisi."),
    ("2. Filter Khusus Per Unit Kerja", "Pimpinan dapat memilih dropdown unit kerja tertentu untuk memeriksa bukti fisik unit tersebut tanpa tercampur."),
    ("3. Kartu Dokumen Eksekutif", "Informasi lengkap: badge unit, kode berkas (DOC-ISI-002), tag SNP, ikon format file (.pdf/.xlsx/.docx), uploader, dan status."),
    ("4. Modal Inspeksi Split-Screen", "• Sisi Kiri: Pratinjau berkas ber-kop resmi SMK Negeri 2 Magelang & checklist bukti fisik.\n• Sisi Kanan: Form catatan koreksi/feedback dan 3 tombol aksi (Sahkan, Minta Revisi, Tolak).")
]

for title, desc in validasi_steps:
    pt = tfl5.add_paragraph()
    pt.text = title
    pt.font.name = "Segoe UI"
    pt.font.size = Pt(10.5)
    pt.font.bold = True
    pt.font.color.rgb = C_NAVY_DARK
    
    pd = tfl5.add_paragraph()
    pd.text = desc
    pd.font.name = "Segoe UI"
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = C_TEXT_MUTED
    pd.space_after = Pt(6)

add_screenshot(slide5, os.path.join(SCREENSHOT_DIR, "01_kasek", "02_validasi_dokumen_kasek.png"), Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.1))

# ==========================================
# SLIDE 6: EVALUASI MUTU & RTL KETUA TPMPS (AKUN TPMPS)
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
add_header(slide6, "TPMPS: Akun 7. UNIT KERJA TPMPS", "Evaluasi Mutu & Rencana Tindak Lanjut (Login sebagai Ketua TPMPS)")

card_l6 = add_card_box(slide6, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.1))
tfl6 = card_l6.text_frame
tfl6.word_wrap = True
tfl6.margin_left = tfl6.margin_right = tfl6.margin_top = Inches(0.25)

p = tfl6.paragraphs[0]
p.text = "TUGAS AUDIT KETUA TPMPS"
p.font.name = "Segoe UI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_AMBER
p.space_after = Pt(10)

tpmps_items = [
    ("Audit Bukti Fisik 14 Unit Kerja", "Memverifikasi kesesuaian dokumen bukti fisik dengan instrumen Standar Nasional Pendidikan."),
    ("Pengisian Skor Evaluasi Mutu", "Menilai skor capaian (skala 1-4) didukung telaah fakta dokumen yang diunggah."),
    ("Penyusunan Catatan Revisi", "Memberikan umpan balik langsung kepada koordinator unit kerja untuk perbaikan."),
    ("Penyusunan RTL (Rencana Tindak Lanjut)", "Menetapkan aksi solutif, target waktu penyelesaian, dan penugasan unit penanggung jawab.")
]

for title, desc in tpmps_items:
    pt = tfl6.add_paragraph()
    pt.text = f"• {title}"
    pt.font.name = "Segoe UI"
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = C_NAVY_DARK
    
    pd = tfl6.add_paragraph()
    pd.text = f"  {desc}"
    pd.font.name = "Segoe UI"
    pd.font.size = Pt(10)
    pd.font.color.rgb = C_TEXT_MUTED
    pd.space_after = Pt(8)

add_screenshot(slide6, os.path.join(SCREENSHOT_DIR, "07_tpmps", "03_evaluasi_mutu_tpmps.png"), Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.1))

# ==========================================
# 15 UNIT SLIDES (LOGGED IN AS EACH UNIT ACCOUNT)
# ==========================================
official_15_units = [
    {
        "unit_id": "1. KASEK",
        "title": "1. KASEK (Kepala Sekolah)",
        "pic": "Drs. H. Mulyono, M.Pd.",
        "email": "kasek@smkn2magelang.sch.id",
        "snp": "Mengesahkan Seluruh 8 Standar Nasional Pendidikan (SNP)",
        "ngapain": [
            "Melakukan validasi dan pengesahan akhir seluruh dokumen bukti fisik yang diunggah oleh 14 unit kerja.",
            "Memantau capaian rapor mutu dan ketercapaian 8 SNP sekolah melalui Executive Dashboard.",
            "Menandatangani dan mengesahkan Laporan Evaluasi Diri Sekolah (EDS) secara resmi.",
            "Mengawasi tindak lanjut rekomendasi perbaikan mutu (RTL) bersama Ketua TPMPS."
        ],
        "bisa_apa": [
            "Executive Dashboard real-time & radar chart capaian 8 SNP.",
            "Fitur Validasi Dokumen: Preview berkas, setujui/sahkan, atau tolak dengan catatan revisi.",
            "Modul Kelola Unit: Monitoring keaktifan unggah dan progres 15 unit kerja.",
            "Modul Laporan EDS: Approval dan penandatanganan legalitas laporan mutu sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "01_kasek", "01_dashboard_kasek.png")
    },
    {
        "unit_id": "2. UNIT KERJA WKS 1",
        "title": "2. UNIT KERJA WKS 1 (Kurikulum)",
        "pic": "Dra. Sri Wahyuni, M.Pd.",
        "email": "wks1@smkn2magelang.sch.id",
        "snp": "Standar Isi, Standar Proses, Standar Penilaian, Standar Kompetensi Lulusan (SKL)",
        "ngapain": [
            "Mengunggah dokumen Kurikulum Operasional Satuan Pendidikan (KOSP) dan Kalender Akademik.",
            "Mengelola perangkat ajar: Alur Tujuan Pembelajaran (ATP), Modul Ajar, dan jadwal KBM guru.",
            "Mengarsipkan bukti pelaksanaan asesmen (Formatif, Sumatif, PTS/PAS, dan UKK).",
            "Menyiapkan eviden supervisi akademik guru dan pembagian jam mengajar (minimal 24 JP)."
        ],
        "bisa_apa": [
            "Upload & kelola bukti fisik kurikulum ke folder standar bersangkutan.",
            "Melihat status verifikasi berkas oleh TPMPS & Kepala Sekolah (Menunggu, Disahkan, Revisi).",
            "Membaca catatan revisi dari pimpinan dan mengunggah dokumen perbaikan.",
            "Memantau persentase kelengkapan dokumen kurikulum secara mandiri."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "02_wks1", "01_drive_wks_1.png")
    },
    {
        "unit_id": "3. UNIT KERJA WKS 2",
        "title": "3. UNIT KERJA WKS 2 (Kesiswaan)",
        "pic": "Bambang Sutrisno, S.Pd.",
        "email": "wks2@smkn2magelang.sch.id",
        "snp": "Standar Kompetensi Lulusan (SKL - Karakter & Sikap), Standar Pengelolaan (Kesiswaan)",
        "ngapain": [
            "Mengunggah dokumen tata tertib sekolah, buku catatan kedisiplinan, dan program pembiasaan karakter.",
            "Mengarsipkan data prestasi siswa lomba akademik/non-akademik tingkat kota/provinsi/nasional.",
            "Mengelola portofolio kegiatan OSIS, MPK, Pramuka, Paskibra, PMR, dan ekstrakurikuler lainnya.",
            "Mendokumentasikan pelaksanaan Masa Pengenalan Lingkungan Sekolah (MPLS) dan beasiswa siswa."
        ],
        "bisa_apa": [
            "Upload sertifikat kejuaraan, foto kegiatan pembinaan karakter, dan laporan kesiswaan.",
            "Melacak status pengesahan dokumen kesiswaan oleh pimpinan.",
            "Menanggapi feedback/catatan revisi dari TPMPS terkait kelengkapan bukti kegiatan.",
            "Memantau progres pemenuhan instrumen karakter dan kesiswaan."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "03_wks2", "01_drive_wks_2.png")
    },
    {
        "unit_id": "4. UNIT KERJA WKS 3",
        "title": "4. UNIT KERJA WKS 3 (Sarana & Prasarana)",
        "pic": "Ir. Agus Haryanto, M.T.",
        "email": "wks3@smkn2magelang.sch.id",
        "snp": "Standar Sarana dan Prasarana (Sarpras)",
        "ngapain": [
            "Mengunggah buku inventarisasi aset, tanah, denah gedung, ruang kelas, laboratorium, dan bengkel.",
            "Mendokumentasikan jadwal dan kartu pemeliharaan berkala sarana dan peralatan praktik.",
            "Mengarsipkan bukti kelaikan sarana K3 (APAR, jalur evakuasi, izin genset, instalasi listrik).",
            "Menyusun usulan kebutuhan belanja modal sarpras (RKAS) dan laporan perbaikan fasilitas."
        ],
        "bisa_apa": [
            "Upload berkas inventaris, sertifikat kelaikan, dan laporan pemeliharaan sarpras.",
            "Melihat status audit kelayakan sarana oleh tim verifikator TPMPS.",
            "Memperbarui berkas sesuai rekomendasi RTL pimpinan terkait fasilitas sekolah.",
            "Memantau skor pemenuhan Standar Sarpras di dashboard unit."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "04_wks3", "01_drive_wks_3.png")
    },
    {
        "unit_id": "5. UNIT KERJA WKS 4",
        "title": "5. UNIT KERJA WKS 4 (Hubin / Humas)",
        "pic": "Drs. Hendro Wibowo",
        "email": "wks4@smkn2magelang.sch.id",
        "snp": "Standar Pengelolaan, Standar Proses (PKL & Kerjasama Industri)",
        "ngapain": [
            "Mengunggah naskah MoU / Perjanjian Kerja Sama (PKS) dengan Dunia Usaha & Dunia Industri (DUDI).",
            "Mengarsipkan pedoman, jurnal bimbingan, dan evaluasi pelaksanaan Praktik Kerja Lapangan (PKL).",
            "Mendokumentasikan program Guru Tamu dari praktisi industri dan magang guru di industri.",
            "Mengarsipkan dokumen sinkronisasi kurikulum berbasis kebutuhan industri dan kelas industri."
        ],
        "bisa_apa": [
            "Upload dokumen legalitas MoU/PKS kemitraan industri beserta masa berlakunya.",
            "Memantau approval berkas kemitraan oleh Kepala Sekolah.",
            "Menerima umpan balik TPMPS mengenai rasio kemitraan DUDI tiap program keahlian.",
            "Meninjau rekapitulasi keterlibatan industri dalam kurikulum sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "05_wks4", "01_drive_wks_4.png")
    },
    {
        "unit_id": "6. UNIT KERJA K3",
        "title": "6. UNIT KERJA K3 (Ketua Kompetensi Keahlian)",
        "pic": "Eko Prasetyo, S.Kom., M.Cs.",
        "email": "k3@smkn2magelang.sch.id",
        "snp": "Standar Isi Kejuruan, Standar Proses Kejuruan, Standar Penilaian Kejuruan, Sarpras Praktik",
        "ngapain": [
            "Mengoordinasikan penyusunan kurikulum operasional konsentrasi keahlian dan jobsheet praktikum.",
            "Mengunggah perangkat Uji Kompetensi Keahlian (UKK) skema LSP-P1 / BNSP / Industri.",
            "Mendokumentasikan portofolio karya/proyek unggulan kejuruan berbasis pesanan nyata (PBL).",
            "Memastikan kesiapan alat bengkel/lab kejuruan sesuai standar sertifikasi profesi."
        ],
        "bisa_apa": [
            "Upload dokumen instrumen kejuruan, rubrik penilaian UKK, dan sertifikasi keahlian.",
            "Melacak verifikasi berkas kesiapan asesmen kejuruan oleh TPMPS.",
            "Merespon catatan revisi reviewer terkait relevansi kurikulum kejuruan dengan standar KKNI.",
            "Memantau kepatuhan standar pembelajaran vokasi di tiap kompetensi keahlian."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "06_k3", "01_drive_k3.png")
    },
    {
        "unit_id": "7. UNIT KERJA TPMPS",
        "title": "7. UNIT KERJA TPMPS (Tim Penjaminan Mutu)",
        "pic": "Dra. Hj. Siti Fatimah, M.M.",
        "email": "tpmps@smkn2magelang.sch.id",
        "snp": "Seluruh 8 Standar Nasional Pendidikan (Audit Mutu Internal, EDS & RTL)",
        "ngapain": [
            "Mengaudit dan memverifikasi kelayakan bukti fisik yang diunggah oleh seluruh 14 unit kerja lainnya.",
            "Mengisi instrumen skor evaluasi mutu (skala 1-4) per butir SNP berdasarkan fakta berkas.",
            "Memberikan feedback / catatan revisi langsung kepada unit kerja jika eviden belum lengkap.",
            "Menyusun dokumen Rencana Tindak Lanjut (RTL) mutu dan analisis kesenjangan (gap analysis)."
        ],
        "bisa_apa": [
            "Modul Validasi: Cek antrean dokumen masuk dari semua unit, preview berkas, sahkan atau minta revisi.",
            "Modul Evaluasi Mutu: Input skor evaluasi per butir standar SNP dan catatan telaah bukti.",
            "Modul RTL Mutu: Formulasi rekomendasi tindak lanjut, tentukan timeline & unit penanggung jawab.",
            "Dashboard TPMPS: Pantauan agregat capaian mutu sekolah sebelum dilaporkan ke Kepala Sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "07_tpmps", "01_dashboard_tpmps.png")
    },
    {
        "unit_id": "8. UNIT KERJA RENBANG",
        "title": "8. UNIT KERJA RENBANG (Perencanaan & Pengembangan)",
        "pic": "Drs. Supriyadi, M.M.",
        "email": "renbang@smkn2magelang.sch.id",
        "snp": "Standar Pengelolaan, Standar Pembiayaan (RKS, RKAS, Rencana Strategis)",
        "ngapain": [
            "Mengunggah dokumen Rencana Kerja Sekolah (RKS 4 Tahunan) dan Rencana Kerja Anggaran Sekolah (RKAS Tahunan).",
            "Menyusun Rencana Pengembangan Sekolah (RPS) berbasis data Rapor Pendidikan Kemendikbud.",
            "Mengarsipkan usulan program prioritas, roadmap pengembangan SMK Pusat Keunggulan, dan diversifikasi program.",
            "Mengkaji kelayakan pengembangan fasilitas baru, teknologi ajar, dan kemitraan strategis."
        ],
        "bisa_apa": [
            "Upload berkas perencanaan strategis (RKS, RKAS, RPS, Roadmap).",
            "Menyelaraskan program perencanaan dengan rekomendasi RTL dari TPMPS.",
            "Memantau persetujuan dokumen perencanaan oleh Kepala Sekolah.",
            "Memantau ketercapaian target milestone pengembangan jangka menengah dan panjang sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "08_renbang", "01_drive_renbang.png")
    },
    {
        "unit_id": "9. UNIT KERJA KATU",
        "title": "9. UNIT KERJA KATU (Kepala Tata Usaha)",
        "pic": "Nurul Hidayati, S.Sos.",
        "email": "katu@smkn2magelang.sch.id",
        "snp": "Standar Pendidik & Tenaga Kependidikan (PTK), Standar Pembiayaan, Standar Pengelolaan",
        "ngapain": [
            "Mengunggah database profil pendidik & tenaga kependidikan (ijazah, Serdik, SK Kenaikan Pangkat/Gaji Berkala).",
            "Mengarsipkan dokumen pertanggungjawaban keuangan (LPJ BOS, BOP, komite) dan buku kas umum.",
            "Mengelola buku induk register siswa, arsip persuratan dinas keluar/masuk, dan legalitas kelembagaan.",
            "Mendokumentasikan Sasaran Kinerja Pegawai (SKP) dan Penilaian Kinerja Guru/Staf (PKG/PKP)."
        ],
        "bisa_apa": [
            "Upload berkas portofolio PTK, LPJ keuangan, dan arsip persuratan dinas.",
            "Melacak status verifikasi dokumen tata usaha oleh tim penjamin mutu.",
            "Menerima catatan audit administrasi dari TPMPS dan Kepala Sekolah.",
            "Memantau indikator kepatuhan rasio PTK dan akuntabilitas keuangan sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "09_katu", "01_drive_katu.png")
    },
    {
        "unit_id": "10. UNIT KERJA KALAB",
        "title": "10. UNIT KERJA KALAB (Kepala Laboratorium)",
        "pic": "Supriyanto, A.Md.",
        "email": "kalab@smkn2magelang.sch.id",
        "snp": "Standar Sarana & Prasarana (Laboratorium), Standar Pengelolaan Lab",
        "ngapain": [
            "Mengunggah Standar Operasional Prosedur (SOP) keselamatan kerja laboratorium dan tata tertib pemakaian.",
            "Mendokumentasikan jadwal penggunaan laboratorium/bengkel untuk KBM dan ujian praktik.",
            "Mengarsipkan buku logbook penggunaan harian dan kartu riwayat pemeliharaan hardware/software/alat.",
            "Menyusun daftar inventaris peralatan lab, spesifikasi PC/alat kerja, dan pengajuan kebutuhan bahan habis pakai."
        ],
        "bisa_apa": [
            "Upload SOP lab, kartu perawatan berkala, dan jadwal pemakaian ruangan lab.",
            "Meninjau persetujuan SOP oleh pimpinan dan catatan evaluasi TPMPS.",
            "Mengunggah bukti tindak lanjut perbaikan alat yang rusak sesuai rekomendasi.",
            "Memantau tingkat kelaikan operasional seluruh fasilitas laboratorium sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "10_kalab", "01_drive_kalab.png")
    },
    {
        "unit_id": "11. UNIT KERJA PERPUSTAKAAN",
        "title": "11. UNIT KERJA PERPUSTAKAAN (Perpustakaan Sekolah)",
        "pic": "Tri Utami, S.I.Pust.",
        "email": "perpustakaan@smkn2magelang.sch.id",
        "snp": "Standar Sarana Perpustakaan, Standar Pengelolaan, Standar Proses (Literasi)",
        "ngapain": [
            "Mengunggah master inventaris koleksi buku teks pelajaran, referensi umum, dan e-book perpustakaan.",
            "Mendokumentasikan statistik kunjungan pemustaka, sirkulasi peminjaman buku, dan kepemilikan kartu anggota.",
            "Mengarsipkan bukti pelaksanaan Gerakan Literasi Sekolah (GLS), lomba literasi, dan pojok baca kelas.",
            "Menyusun laporan tata kelola otomasi perpustakaan (Inlislite/Senayan) dan pengadaan buku baru."
        ],
        "bisa_apa": [
            "Upload katalog buku, laporan sirkulasi bulanan, dan eviden kegiatan literasi sekolah.",
            "Memantau status pengesahan berkas perpustakaan oleh Kepala Sekolah.",
            "Menerima umpan balik dari TPMPS terkait pemenuhan rasio buku per peserta didik.",
            "Memantau ketercapaian Standar Nasional Perpustakaan (SNP)."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "11_perpustakaan", "01_drive_perpustakaan.png")
    },
    {
        "unit_id": "12. UNIT KERJA NASWIL",
        "title": "12. UNIT KERJA NASWIL (Nasionalisme & Bina Wilayah)",
        "pic": "Drs. H. Mulyadi, M.Pd.",
        "email": "naswil@smkn2magelang.sch.id",
        "snp": "Standar Kompetensi Lulusan (SKL - Nilai Karakter Pancasila & Kebangsaan), Standar Pengelolaan",
        "ngapain": [
            "Mengunggah bukti pelaksanaan upacara bendera, apel kebangsaan, dan peringatan hari besar nasional (PHBN).",
            "Mendokumentasikan kegiatan Bela Negara, pembinaan wawasan nusantara, dan kemitraan dengan Koramil/Polsek/BNN.",
            "Mengarsipkan laporan program sekolah Adiwiyata, bakti sosial masyarakat sekitar, dan bina lingkungan kewilayahan.",
            "Menyusun portofolio pencegahan intoleransi, anti-perundungan (anti-bullying), dan radikalisme di sekolah."
        ],
        "bisa_apa": [
            "Upload laporan kegiatan kebangsaan, foto aksi sosial kemasyarakatan, dan MoU dengan instansi kewilayahan.",
            "Memantau status verifikasi dokumen pembiasaan nasionalisme oleh TPMPS.",
            "Merespon masukan perbaikan program pembinaan karakter wawasan kebangsaan.",
            "Memantau pemenuhan profil pelajar Pancasila dan indikator bina lingkungan sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "12_naswil", "01_drive_naswil.png")
    },
    {
        "unit_id": "13. UNIT KERJA BK",
        "title": "13. UNIT KERJA BK (Bimbingan Konseling)",
        "pic": "Dra. Endang Sulastri",
        "email": "bk@smkn2magelang.sch.id",
        "snp": "Standar Proses (Layanan BK), Standar Kompetensi Lulusan (Pengembangan Diri & Karir)",
        "ngapain": [
            "Mengunggah program kerja tahunan layanan bimbingan: pribadi, sosial, belajar, dan karir.",
            "Mendokumentasikan asesmen diagnostik non-kognitif awal tahun untuk pemetaan gaya belajar dan potensi siswa.",
            "Mengarsipkan rekapitulasi layanan konseling perorangan, bimbingan kelompok, serta penanganan kasus siswa.",
            "Menyusun laporan peminatan jurusan, bimbingan studi lanjut ke perguruan tinggi, dan pembinaan karir kerja."
        ],
        "bisa_apa": [
            "Upload dokumen program kerja BK, jadwal layanan, dan rekapitulasi laporan umum.",
            "Memantau persetujuan program kerja BK oleh Kepala Sekolah.",
            "Menerima arahan peningkatan mutu layanan konseling dari TPMPS.",
            "Memantau indikator ketercapaian layanan pengembangan karakter dan kepribadian siswa."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "13_bk", "01_drive_bk.png")
    },
    {
        "unit_id": "14. UNIT KERJA BKK",
        "title": "14. UNIT KERJA BKK (Bursa Kerja Khusus)",
        "pic": "Wahyu Nugroho, S.Pd.",
        "email": "bkk@smkn2magelang.sch.id",
        "snp": "Standar Kompetensi Lulusan (SKL - Keterserapan Lulusan di DUDI / Tracer Study)",
        "ngapain": [
            "Mengunggah rekapitulasi data Tracer Study: Bekerja, Melanjutkan Pendidikan, Wirausaha (BMW).",
            "Mendokumentasikan pelaksanaan Bursa Kerja Khusus (Job Fair), rekrutmen langsung kampus, dan seleksi industri.",
            "Mengarsipkan surat penempatan kerja lulusan dan data serapan kerja kurang dari 6 bulan pasca lulus.",
            "Menyusun laporan kerjasama dengan Dinas Tenaga Kerja dan forum BKK SMK tingkat wilayah."
        ],
        "bisa_apa": [
            "Upload database penelusuran tamatan, bukti rekrutmen DUDI, dan foto kegiatan job fair.",
            "Memantau validasi data serapan lulusan oleh Kepala Sekolah.",
            "Menerima rekomendasi tindak lanjut dari TPMPS mengenai strategi peningkatan persentase serapan alumni.",
            "Menyajikan statistik keterserapan tamatan secara transparan untuk akreditasi dan rapor mutu sekolah."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "14_bkk", "01_drive_bkk.png")
    },
    {
        "unit_id": "15. UNIT KERJA UPS",
        "title": "15. UNIT KERJA UPS (Unit Produksi Sekolah / TEFA)",
        "pic": "Anwar Sadat, S.T.",
        "email": "ups@smkn2magelang.sch.id",
        "snp": "Standar Proses, Standar Pengelolaan, Standar Pembiayaan (Pembelajaran Berbasis Industri & Pendapatan Jasa/Produk)",
        "ngapain": [
            "Mengunggah pedoman operasional Unit Produksi Sekolah (UPS) dan alur kerja Teaching Factory (TEFA).",
            "Mendokumentasikan pesanan pekerjaan nyata (job order) dari masyarakat/industri yang dikerjakan siswa.",
            "Mengarsipkan laporan omzet penjualan produk/jasa, neraca laba-rugi UPS, dan kontribusi terhadap sekolah.",
            "Menyusun portofolio produk barang/jasa unggulan karya siswa yang telah dipasarkan."
        ],
        "bisa_apa": [
            "Upload SOP produksi, katalog produk/jasa UPS, dan rekapitulasi laporan keuangan unit produksi.",
            "Memantau pengesahan dokumen kepatuhan mutu TEFA oleh Kepala Sekolah.",
            "Menerima catatan telaah dari TPMPS terkait integrasi kurikulum kejuruan dengan lini produksi UPS.",
            "Menampilkan kinerja profesionalisme SMK Negeri 2 Magelang sebagai lembaga vokasi berbasis produk nyata."
        ],
        "img": os.path.join(SCREENSHOT_DIR, "15_ups", "01_drive_ups.png")
    }
]

for u in official_15_units:
    slide_u = prs.slides.add_slide(blank_layout)
    add_header(slide_u, f"Portofolio Unit Kerja: {u['unit_id']}", f"{u['title']} (Login Akun Unit)")
    
    card_u = add_card_box(slide_u, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.1))
    tfu = card_u.text_frame
    tfu.word_wrap = True
    tfu.margin_left = tfu.margin_right = tfu.margin_top = Inches(0.2)
    
    p_pic = tfu.paragraphs[0]
    p_pic.text = f"Akun / PIC: {u['pic']}"
    p_pic.font.name = "Segoe UI"
    p_pic.font.size = Pt(10.5)
    p_pic.font.bold = True
    p_pic.font.color.rgb = C_BLUE_MAIN
    p_pic.space_after = Pt(1)
    
    p_mail = tfu.add_paragraph()
    p_mail.text = f"Email Login: {u['email']}"
    p_mail.font.name = "Segoe UI"
    p_mail.font.size = Pt(9)
    p_mail.font.color.rgb = C_AMBER
    p_mail.space_after = Pt(2)
    
    p_snp = tfu.add_paragraph()
    p_snp.text = f"SNP Diampu: {u['snp']}"
    p_snp.font.name = "Segoe UI"
    p_snp.font.size = Pt(8.5)
    p_snp.font.italic = True
    p_snp.font.color.rgb = C_TEXT_MUTED
    p_snp.space_after = Pt(6)
    
    # Section 1: Ngapain Aja?
    p_sec1 = tfu.add_paragraph()
    p_sec1.text = "TUGAS POKOK & BUKTI FISIK (Ngapain Aja?):"
    p_sec1.font.name = "Segoe UI"
    p_sec1.font.size = Pt(10)
    p_sec1.font.bold = True
    p_sec1.font.color.rgb = C_NAVY_DARK
    p_sec1.space_after = Pt(2)
    
    for item in u["ngapain"]:
        pi = tfu.add_paragraph()
        pi.text = f"• {item}"
        pi.font.name = "Segoe UI"
        pi.font.size = Pt(8.5)
        pi.font.color.rgb = C_TEXT_DARK
        pi.space_after = Pt(2.5)
        
    p_gap = tfu.add_paragraph()
    p_gap.space_after = Pt(3)
    
    # Section 2: Bisa Apa Aja di Sistem?
    p_sec2 = tfu.add_paragraph()
    p_sec2.text = "FITUR & AKSES SISTEM SINTESA (Bisa Apa Aja?):"
    p_sec2.font.name = "Segoe UI"
    p_sec2.font.size = Pt(10)
    p_sec2.font.bold = True
    p_sec2.font.color.rgb = C_AMBER
    p_sec2.space_after = Pt(2)
    
    for item in u["bisa_apa"]:
        pi = tfu.add_paragraph()
        pi.text = f"✓ {item}"
        pi.font.name = "Segoe UI"
        pi.font.size = Pt(8.5)
        pi.font.color.rgb = C_TEXT_DARK
        pi.space_after = Pt(2.5)
        
    add_screenshot(slide_u, u["img"], Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.1))

# ==========================================
# SLIDE 22: ALUR SIKLUS PPEPP TERPADU
# ==========================================
slide22 = prs.slides.add_slide(blank_layout)
add_header(slide22, "Siklus Penjaminan Mutu Internal", "Alur Kerja End-to-End PPEPP & Siklus Validasi SINTESA")

ppepp_steps = [
    ("1. Penetapan (P)", "Pimpinan & TPMPS menetapkan target 8 SNP, indikator mutu, butir instrumen, dan penugasan 15 unit kerja.", C_NAVY_MED),
    ("2. Pelaksanaan (P)", "15 Unit Kerja mengunggah berkas bukti fisik, dokumen portofolio, dan laporan kegiatan ke Drive Unit masing-masing.", C_BLUE_MAIN),
    ("3. Evaluasi (E)", "Ketua TPMPS & Tim mengaudit keabsahan berkas di Modul Validasi, menginput skor evaluasi, dan memberikan catatan koreksi.", C_AMBER),
    ("4. Pengendalian (P)", "Kepala Sekolah mengesahkan dokumen yang valid atau menolak untuk revisi. Menyusun Laporan EDS resmi dan penandatanganan legal.", C_EMERALD),
    ("5. Peningkatan (P)", "Penyusunan Rencana Tindak Lanjut (RTL) berbasis kesenjangan mutu untuk peningkatan standar di periode berikutnya.", RGBColor(124, 58, 237))
]

for idx, (step_title, step_desc, step_color) in enumerate(ppepp_steps):
    y_pos = Inches(1.8 + idx * 1.0)
    box = add_card_box(slide22, Inches(0.8), y_pos, Inches(11.733), Inches(0.85))
    tf_step = box.text_frame
    tf_step.word_wrap = True
    tf_step.margin_left = tf_step.margin_right = Inches(0.2)
    tf_step.margin_top = Inches(0.12)
    
    p1 = tf_step.paragraphs[0]
    p1.text = step_title
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = step_color
    p1.space_after = Pt(2)
    
    p2 = tf_step.add_paragraph()
    p2.text = step_desc
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(10)
    p2.font.color.rgb = C_NAVY_DARK

# ==========================================
# SLIDE 23: KESIMPULAN & PENGESAHAN
# ==========================================
slide23 = prs.slides.add_slide(blank_layout)
add_header(slide23, "Penutup & Lembar Pengesahan", "Komitmen Bersama Penjaminan Mutu Berkelanjutan")

card_l23 = add_card_box(slide23, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
tf23 = card_l23.text_frame
tf23.word_wrap = True
tf23.margin_left = tf23.margin_right = tf23.margin_top = Inches(0.3)

p = tf23.paragraphs[0]
p.text = "DAMPAK IMPLEMENTASI SINTESA TPMPS"
p.font.name = "Segoe UI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_BLUE_MAIN
p.space_after = Pt(14)

impacts = [
    ("Akuntabilitas Terukur", "Setiap bukti fisik dari 15 unit kerja terdokumentasi rapi, terverifikasi berjenjang, dan bebas dari risiko kehilangan arsip."),
    ("Transparansi Kinerja Unit", "Pimpinan dapat memantau produktivitas dan kepatuhan 15 unit kerja secara real-time dari satu dashboard terpadu."),
    ("Efisiensi Akreditasi Sekolah", "Seluruh eviden siap saji untuk kebutuhan asesmen BAN-PDM, audit ISO, maupun evaluasi dinas pendidikan."),
    ("Budaya Mutu Berkelanjutan", "Siklus PPEPP bertransformasi dari sekadar kewajiban administratif menjadi budaya kerja nyata di SMK Negeri 2 Magelang.")
]

for title, desc in impacts:
    pt = tf23.add_paragraph()
    pt.text = f"✔ {title}"
    pt.font.name = "Segoe UI"
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = C_NAVY_DARK
    
    pd = tf23.add_paragraph()
    pd.text = f"   {desc}"
    pd.font.name = "Segoe UI"
    pd.font.size = Pt(10)
    pd.font.color.rgb = C_TEXT_MUTED
    pd.space_after = Pt(10)

card_r23 = add_card_box(slide23, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1))
tfr23 = card_r23.text_frame
tfr23.word_wrap = True
tfr23.margin_left = tfr23.margin_right = tfr23.margin_top = Inches(0.3)

ps_title = tfr23.paragraphs[0]
ps_title.text = "LEMBAR PENGESAHAN DOKUMEN SISTEM"
ps_title.font.name = "Segoe UI"
ps_title.font.size = Pt(13)
ps_title.font.bold = True
ps_title.font.color.rgb = C_NAVY_DARK
ps_title.alignment = PP_ALIGN.CENTER
ps_title.space_after = Pt(4)

ps_sub = tfr23.add_paragraph()
ps_sub.text = "Magelang, September 2026"
ps_sub.font.name = "Segoe UI"
ps_sub.font.size = Pt(10)
ps_sub.font.color.rgb = C_TEXT_MUTED
ps_sub.alignment = PP_ALIGN.CENTER
ps_sub.space_after = Pt(30)

sig_text = tfr23.add_paragraph()
sig_text.text = "Mengetahui,\nKepala SMK Negeri 2 Magelang\n\n\n\n\n( Drs. H. Mulyono, M.Pd. )\nNIP. 196803121992031004"
sig_text.font.name = "Segoe UI"
sig_text.font.size = Pt(10.5)
sig_text.font.color.rgb = C_NAVY_DARK
sig_text.alignment = PP_ALIGN.CENTER
sig_text.space_after = Pt(25)

sig_tpmps = tfr23.add_paragraph()
sig_tpmps.text = "Disetujui oleh,\nKetua TPMPS SMK Negeri 2 Magelang\n\n\n\n\n( Dra. Hj. Siti Fatimah, M.M. )\nNIP. 197509182002122001"
sig_tpmps.font.name = "Segoe UI"
sig_tpmps.font.size = Pt(10.5)
sig_tpmps.font.color.rgb = C_NAVY_DARK
sig_tpmps.alignment = PP_ALIGN.CENTER

# Save Output
output_path1 = r"C:\Users\lulus\tpmps-smk\PRESENTASI_SINTESA_TPMPS.pptx"
output_path2 = r"C:\Users\lulus\tpmps-smk\public\PRESENTASI_SINTESA_TPMPS.pptx"

prs.save(output_path1)
prs.save(output_path2)

print(f"NEW Presentation generated successfully using PER-ACCOUNT screenshots at:\n1. {output_path1}\n2. {output_path2}")
